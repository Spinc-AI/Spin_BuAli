"""Report export: RTL detection and output filenames."""
import pytest
from docx import Document

from export import build_filename, looks_rtl, sanitize_filename_part, save_as_docx


class TestRtlDetection:
    @pytest.mark.parametrize("text", [
        "کلیه راست طبیعی است",
        "Right kidney 107 x 44 mm — کلیه راست",  # mixed script still counts
    ])
    def test_arabic_script_is_rtl(self, text):
        assert looks_rtl(text)

    @pytest.mark.parametrize("text", ["Right kidney measures 107 x 44 mm.", "", "123 mm"])
    def test_latin_and_empty_are_not(self, text):
        assert not looks_rtl(text)


class TestFilenames:
    @pytest.mark.parametrize("model, expected", [
        ("openai:gpt-4o-mini", "gpt-4o-mini"),
        ("gemini:gemini-2.5-pro", "gemini-2.5-pro"),
        ("aya-expanse-8b", "aya-expanse-8b"),
        ("", ""),
    ])
    def test_prefixes_are_dropped(self, model, expected):
        assert sanitize_filename_part(model) == expected

    def test_path_characters_are_replaced(self):
        assert "/" not in sanitize_filename_part("some/model:v1")
        assert sanitize_filename_part('a"b<c>d') == "a-b-c-d"

    def test_name_combines_recording_and_model(self):
        assert build_filename("/audio/DPM89130.MP3", ["gemini:gemini-2.5-pro"]) == \
            "DPM89130_gemini-2.5-pro.docx"

    def test_falls_back_when_there_is_no_recording(self):
        assert build_filename(None, ["aya-expanse-8b"]) == "report_aya-expanse-8b.docx"

    def test_empty_model_parts_are_skipped(self):
        assert build_filename("/audio/clip.wav", ["", None]) == "clip.docx"


class TestDocxOutput:
    def test_writes_a_readable_document(self, tmp_path):
        path = tmp_path / "report.docx"
        save_as_docx("Right kidney 107 mm.\nکلیه راست طبیعی است", str(path))

        paragraphs = Document(str(path)).paragraphs
        assert [p.text for p in paragraphs] == ["Right kidney 107 mm.", "کلیه راست طبیعی است"]

    def test_only_the_rtl_paragraph_is_marked(self, tmp_path):
        path = tmp_path / "report.docx"
        save_as_docx("Latin line\nخط فارسی", str(path))

        latin, persian = Document(str(path)).paragraphs
        assert "bidi" not in latin._p.xml
        assert "bidi" in persian._p.xml
